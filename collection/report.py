"""集計結果を Excel / CSV に出力するモジュール。"""
from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_FIXED_WIDTH = {
    "コード": 8,
    "銘柄名": 26,
    "33業種区分": 16,
    "規模区分": 14,
    "ステータス": 14,
    "決算短信表題": 40,
    "開示日": 12,
    "PDFファイル": 34,
    "抽出文字数": 10,
    "該当有無": 9,
    "ヒットワード数": 12,
    "合計ヒット数": 12,
    "代表的な出現箇所": 60,
    "備考": 30,
}


def write_csv(rows: list[dict], path: Path) -> None:
    """Excel で開ける UTF-8(BOM付き) CSV を書き出す。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(rows: list[dict], summary: dict, path: Path) -> None:
    """一覧シートとサマリシートを持つ Excel ブックを書き出す。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    headers = list(rows[0].keys())

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ワード集計"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    for column_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column_index, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        letter = get_column_letter(column_index)
        sheet.column_dimensions[letter].width = _FIXED_WIDTH.get(header, 11)

    for row_index, row in enumerate(rows, start=2):
        for column_index, header in enumerate(headers, start=1):
            sheet.cell(row=row_index, column=column_index, value=row[header])

    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    )

    summary_sheet = workbook.create_sheet("サマリ")
    summary_sheet.column_dimensions["A"].width = 28
    summary_sheet.column_dimensions["B"].width = 16
    line = 1
    for key, value in summary.items():
        if isinstance(value, dict):
            continue
        summary_sheet.cell(row=line, column=1, value=key).font = Font(bold=True)
        summary_sheet.cell(row=line, column=2, value=value)
        line += 1

    line += 1
    summary_sheet.cell(row=line, column=1, value="ワード別 該当社数").font = Font(bold=True)
    line += 1
    for word, count in summary["ワード別該当社数"].items():
        summary_sheet.cell(row=line, column=1, value=word)
        summary_sheet.cell(row=line, column=2, value=count)
        line += 1

    workbook.save(path)
